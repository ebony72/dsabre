#!/usr/bin/env python3
"""Build the general-audience dSABRE talk (22 slides) with python-pptx.

Run:  python3 build_deck.py
Out:  dSABRE_talk.pptx  (16:9, 13.333 x 7.5 in)

Figures in img/ are produced by ../paper/fig/*.tex (recompiled standalone)
and by the chart scripts recorded in README_talk.md.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

_HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(_HERE, "img")

# ── palette ──────────────────────────────────────────────────────────────────
INK        = "0E1730"   # deep navy — dark slides, headings
INK_SOFT   = "1C2743"
TEAL       = "12867E"   # the accent, matches the figures' comm ports
TEAL_DARK  = "0C6259"
TEAL_TINT  = "E7F2F0"
CORAL      = "C0453B"
CORAL_TINT = "FBEDEB"
GREY       = "58melt"   # placeholder, replaced below
GREY       = "586178"
GREY_LT    = "8A93A6"
PANEL      = "F4F6F9"
WHITE      = "FFFFFF"

HEAD = "Cambria"
BODY = "Calibri"

SW, SH = 13.333, 7.5
M      = 0.62
CW     = SW - 2 * M

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# ── low-level helpers ────────────────────────────────────────────────────────
def rgb(h):
    return RGBColor.from_string(h)


def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK if dark else WHITE)
    return s


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, first=False):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def run(p, text, size=14, color=INK, bold=False, italic=False,
        font=BODY, spc=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = rgb(color)
    if spc is not None:                       # letter-spacing, 1/100 pt
        r.font._rPr.set("spc", str(int(spc * 100)))
    return r


def write(tf, text, size=14, color=INK, bold=False, italic=False,
          font=BODY, first=False, align=PP_ALIGN.LEFT, space_after=0,
          line_spacing=None, spc=None):
    p = para(tf, first)
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = line_spacing
    run(p, text, size, color, bold, italic, font, spc)
    return p


def bullet(p, char="•", color=TEAL, indent=0.20):
    """Attach a real bullet glyph. Call after all other paragraph props."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent)).emu))
    pPr.set("indent", str(-Emu(Inches(indent)).emu))
    for tag, attrs in (("a:buClr", None),
                       ("a:buFont", {"typeface": "Arial"}),
                       ("a:buChar", {"char": char})):
        el = pPr.makeelement(qn(tag), attrs or {})
        if tag == "a:buClr":
            el.append(pPr.makeelement(qn("a:srgbClr"), {"val": color}))
        pPr.append(el)


def bullets(tf, items, size=14, color=INK, gap=9, char="•",
            bullet_color=TEAL, first=True, line_spacing=1.12, indent=0.20):
    """items: list of str, or (bold_lead, rest) tuples."""
    for i, it in enumerate(items):
        p = para(tf, first and i == 0)
        p.space_after = Pt(gap)
        p.line_spacing = line_spacing
        if isinstance(it, tuple):
            run(p, it[0], size, color, bold=True)
            if it[1]:
                run(p, it[1], size, color)
        else:
            run(p, it, size, color)
        bullet(p, char, bullet_color, indent)


def card(slide, x, y, w, h, fill=PANEL, line=None, radius=0.055):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1.0)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.text_frame.text = ""
    return sh


def circle_num(slide, x, y, d, n, fill=TEAL, color=WHITE, size=16):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, str(n), size, color, bold=True, font=HEAD)
    return sh


def picture(slide, name, x, w, y):
    """Place img/<name>.png at (x, y) with width w; returns bottom edge."""
    path = os.path.join(IMG, name + ".png")
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    return y + pic.height / 914400.0


def head(slide, kicker, title, dark=False):
    tf = textbox(slide, M, 0.40, CW, 0.28)
    write(tf, kicker.upper(), 10.5, TEAL if not dark else "5FCFC2",
          bold=True, first=True, spc=1.4)
    tf = textbox(slide, M, 0.68, CW, 0.80)
    write(tf, title, 30, WHITE if dark else INK, bold=True,
          font=HEAD, first=True)


def footer(slide, n):
    tf = textbox(slide, SW - 1.30, 6.99, 0.68, 0.30)
    write(tf, str(n), 10, GREY_LT, first=True, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def caption(slide, x, y, w, text, size=11, color=GREY, align=PP_ALIGN.LEFT,
            italic=False, h=0.42):
    tf = textbox(slide, x, y, w, h)
    write(tf, text, size, color, italic=italic, first=True, align=align,
          line_spacing=1.10)


# ═════════════════════════════════════════════════════════════════════════════
# 1 — Title
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide(dark=True)

tf = textbox(s, M + 0.30, 1.62, 8.4, 0.32)
write(tf, "IEEE TRANSACTIONS ON CAD  ·  UNDER REVISION", 11, "5FCFC2",
      bold=True, first=True, spc=1.6)

tf = textbox(s, M + 0.30, 2.06, 8.6, 1.15)
write(tf, "dSABRE", 62, WHITE, bold=True, font=HEAD, first=True)

tf = textbox(s, M + 0.30, 3.28, 7.9, 1.55)
write(tf, "Compiling quantum programs for machines\nmade of many small quantum chips",
      23, "C9D2E4", font=HEAD, first=True, line_spacing=1.22)

tf = textbox(s, M + 0.30, 5.30, 7.6, 1.20)
write(tf, "Sanjiang Li", 17, WHITE, bold=True, first=True, space_after=4)
write(tf, "Centre for Quantum Software and Information,\nUniversity of Technology Sydney",
      12.5, "94A0BA", line_spacing=1.18)

# motif: a 2x3 core network echoing the H-grid architecture
ox, oy, cw, ch, gx, gy = 9.55, 2.30, 0.92, 0.92, 1.16, 1.16
centres = {}
for r_ in range(2):
    for c_ in range(3):
        x0, y0 = ox + c_ * gx, oy + r_ * gy
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0),
                                Inches(y0), Inches(cw), Inches(ch))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb("18304E")
        sh.line.color.rgb = rgb("2C6E6A")
        sh.line.width = Pt(1.0)
        sh.shadow.inherit = False
        try:
            sh.adjustments[0] = 0.16
        except Exception:
            pass
        centres[(r_, c_)] = (x0 + cw / 2, y0 + ch / 2)

for a, b in [((0, 0), (0, 1)), ((0, 1), (0, 2)), ((1, 0), (1, 1)),
             ((1, 1), (1, 2)), ((0, 0), (1, 0)), ((0, 2), (1, 2)),
             ((0, 1), (1, 1))]:
    (x1, y1), (x2, y2) = centres[a], centres[b]
    cn = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = rgb("2E8C84")
    cn.line.width = Pt(1.5)

for pos in centres.values():
    d = 0.17
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(pos[0] - d / 2),
                             Inches(pos[1] - d / 2), Inches(d), Inches(d))
    dot.fill.solid()
    dot.fill.fore_color.rgb = rgb(TEAL)
    dot.line.fill.background()
    dot.shadow.inherit = False

caption(s, 9.55, 4.86, 3.3, "six chips, seven entanglement links", 10.5,
        "6B7races" if False else "6B788F", italic=True)

notes(s, """
Good afternoon. This talk is about a compiler — specifically, about one decision a
compiler has to make thousands of times when a quantum program runs on a machine
built from several small quantum chips rather than one big one.

I'm going to spend the first half of the talk on background, because the interesting
part of the problem only makes sense once you know what a teleport actually costs.
No quantum mechanics background is assumed. If you take one thing away: the expensive
resource in a distributed quantum computer is entanglement, and the compiler decides
how much of it you burn.
""")
footer(s, 1)

# ═════════════════════════════════════════════════════════════════════════════
# 2 — Roadmap
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Roadmap", "Three questions")

qs = [("What is a quantum computer, and why can't we just build a big one?",
       "Qubits, circuits, and the reason chips stay small."),
      ("What is teleportation, and why does it cost something?",
       "Entanglement as fuel: slow to make, noisy, and used up."),
      ("How should a compiler\nspend that cost?",
       "This is the research contribution — and it halves the bill.")]

cx, cw_, gap = M, 3.83, 0.31
for i, (q, sub) in enumerate(qs):
    x = cx + i * (cw_ + gap)
    card(s, x, 1.72, cw_, 3.52, fill=PANEL)
    circle_num(s, x + 0.34, 2.06, 0.62, i + 1)
    tf = textbox(s, x + 0.34, 2.86, cw_ - 0.68, 1.66,
                 anchor=MSO_ANCHOR.BOTTOM)
    write(tf, q, 16.5, INK, bold=True, font=HEAD, first=True, line_spacing=1.14)
    tf = textbox(s, x + 0.34, 4.66, cw_ - 0.68, 0.72)
    write(tf, sub, 12.5, GREY, first=True, line_spacing=1.16)

card(s, M, 5.52, CW, 1.02, fill=TEAL_TINT)
tf = textbox(s, M + 0.40, 5.76, CW - 0.80, 0.60, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "The short version:  a cross-chip move is not just an expensive local move, "
          "and pricing it properly is worth about half the entanglement budget.",
      15, TEAL_DARK, bold=True, first=True, line_spacing=1.12)

notes(s, """
Here's the shape of the talk. Roughly the first ten minutes is background — qubits,
circuits, what a compiler for a quantum computer actually does, and then teleportation
in detail. The second half is the actual research.

If you already know quantum computing, the part that will be new starts at the slide
titled "A teleport is not just an expensive SWAP."
""")
footer(s, 2)

# ═════════════════════════════════════════════════════════════════════════════
# 3 — Bits and qubits
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Background  ·  1 of 5", "Bits, and what a qubit does differently")

tf = textbox(s, M, 1.74, 5.55, 4.4)
write(tf, "A classical bit is 0 or 1. That's the whole story, and it is why classical "
          "computers are so reliable: a bit that drifts slightly is snapped back to "
          "the nearest of two values.", 15, INK, first=True, space_after=13,
      line_spacing=1.20)
write(tf, "A qubit is not restricted to those two values, and a group of qubits can "
          "hold correlations that no group of classical bits can reproduce. That is "
          "where the computational power comes from — and also where all the "
          "difficulty comes from.", 15, INK, space_after=13, line_spacing=1.20)
write(tf, "And you cannot peek to check. Looking at a qubit collapses it to a plain "
          "0 or 1, so a compiler gets no feedback from the machine — it has to be "
          "right the first time.", 15, INK, line_spacing=1.20)

card(s, M, 5.30, 5.55, 1.26, fill=CORAL_TINT)
tf = textbox(s, M + 0.34, 5.54, 5.55 - 0.68, 0.80, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "Why this matters for a compiler:  every operation is delicate, so the "
          "fewer operations we run, the better the answer.", 14, "8E322B",
      bold=True, first=True, line_spacing=1.14)

props = [("Superposition",
          "A qubit holds a blend of 0 and 1 at once. Ask it for an answer and it "
          "gives you one of them, with a probability set by the blend."),
         ("Entanglement",
          "Two qubits can share a state that neither one owns on its own. Measuring "
          "one instantly constrains the other, however far apart they are."),
         ("Fragility",
          "Reading a qubit destroys what it held, and an unknown qubit cannot be "
          "copied. Every gate you run adds a little noise.")]

px, pw = 6.55, 6.16
for i, (t, d) in enumerate(props):
    y = 1.74 + i * 1.65
    card(s, px, y, pw, 1.44, fill=WHITE, line="DDE3EA")
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px + 0.32),
                             Inches(y + 0.30), Inches(0.20), Inches(0.20))
    dot.fill.solid()
    dot.fill.fore_color.rgb = rgb(TEAL)
    dot.line.fill.background()
    dot.shadow.inherit = False
    tf = textbox(s, px + 0.66, y + 0.24, pw - 1.00, 0.34)
    write(tf, t, 15.5, INK, bold=True, font=HEAD, first=True)
    tf = textbox(s, px + 0.66, y + 0.62, pw - 1.00, 0.70)
    write(tf, d, 12.5, GREY, first=True, line_spacing=1.15)

notes(s, """
Three properties, and only three, are needed for the rest of the talk.

Superposition is the one everyone has heard of. Entanglement is the one that matters
most here — it is literally the resource we are going to be counting. And fragility is
the reason the whole compilation problem exists: every extra operation degrades the
answer, so a compiler that inserts fewer operations produces a better result on real
hardware.

Note the last point on the right: an unknown qubit cannot be copied. That is the
no-cloning theorem, and it is why moving a qubit between chips is not as simple as
copying a file across a network. Remember it for the teleportation slide.
""")
footer(s, 3)

# ═════════════════════════════════════════════════════════════════════════════
# 4 — Circuits and the DAG
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Background  ·  2 of 5", "A quantum program is a circuit")

tf = textbox(s, M, 1.78, 4.55, 4.5)
bullets(tf, [
    ("Wires are qubits, ", "boxes are gates. Time runs left to right."),
    ("Two kinds of gate are enough. ", "Single-qubit rotations, plus one "
     "two-qubit gate called CX, can express any quantum computation."),
    ("The compiler doesn't see a picture. ", "It sees a dependency graph: gate "
     "g₃ cannot run until g₁ and g₂ are done, because they share qubits."),
    ("The front layer ", "is the set of gates with nothing left to wait for — "
     "what could run right now. Everything in this talk is steered by it."),
], size=14.5, gap=12)

y_img = picture(s, "dag_example", x=5.55, w=7.30, y=2.05)
caption(s, 5.55, y_img + 0.22, 7.30, "The same three-gate program, twice: as a circuit "
        "diagram, and as the dependency graph a router actually works on. The two "
        "blue gates are the front layer.", 11.5, GREY, h=0.72)

notes(s, """
Left: the vocabulary. A quantum circuit is drawn like sheet music — one horizontal
wire per qubit, boxes for operations. Universality means you don't need many kinds of
box: arbitrary one-qubit rotations plus the two-qubit CX gate can build anything.

Right: how the compiler sees it. Strip out the drawing and what remains is a
dependency graph — a DAG. g3 acts on qubits 2 and 3, so it has to wait for the two
gates that touch those qubits.

The front layer, in blue, is the set of gates that are ready now. Every routing
decision in this talk is scored by how much it helps the front layer. Keep that in
mind — it is the one idea the entire algorithm is built around.
""")
footer(s, 4)

# ═════════════════════════════════════════════════════════════════════════════
# 5 — Routing on one chip
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Background  ·  3 of 5", "The catch: qubits can only talk to their neighbours")

tf = textbox(s, M, 1.62, CW, 0.42)
write(tf, "On real hardware a two-qubit gate only works if its two qubits sit next to "
          "each other. Usually they don't.", 15, GREY, first=True)

y_img = picture(s, "routing_onechip", x=2.22, w=8.90, y=2.05)

card(s, M, 6.28, CW, 0.86, fill=TEAL_TINT)
tf = textbox(s, M + 0.40, 6.46, CW - 0.80, 0.52, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "This is the qubit routing problem — and it is pure overhead. "
          "None of those SWAPs is part of the program the user wrote.",
      14.5, TEAL_DARK, bold=True, first=True)

notes(s, """
This is the problem the whole field of qubit routing exists to solve.

A chip has a fixed wiring pattern — here a four-by-four grid. A two-qubit gate is
physically possible only between qubits that are wired together. So when the program
asks for a gate between two qubits that are three hops apart, the compiler has to
physically shuffle them together first, using SWAP operations.

Two things to notice. First, a SWAP is not free: it compiles down to three CX gates,
each of which adds noise. Second, none of this work is in the user's program — it is
pure overhead introduced by the mismatch between the program and the hardware.
Minimising it is the compiler's job.
""")
footer(s, 5)

# ═════════════════════════════════════════════════════════════════════════════
# 6 — SABRE
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Background  ·  4 of 5", "SABRE: the routing heuristic everyone ships")

tf = textbox(s, M, 1.78, 6.05, 4.3)
bullets(tf, [
    ("It is the default in Qiskit, ", "IBM's compiler — so it is what most "
     "quantum programs in the world actually go through."),
    ("The rule is simple. ", "Look at the gates that are ready now. Of all the "
     "SWAPs you could do, pick the one that brings those gates' qubits closest "
     "together."),
    ("With a little lookahead. ", "Also glance at the gates coming up next, at a "
     "discount, so the router doesn't walk itself into a corner."),
    ("It is not optimal, ", "and it does not try to be. It is fast, predictable, "
     "and good enough that production toolchains ship it."),
], size=14.5, gap=12)

card(s, 7.10, 1.82, 5.61, 3.06, fill=INK)
tf = textbox(s, 7.48, 2.08, 4.90, 0.32)
write(tf, "THE SCORE, IN WORDS", 10.5, "5FCFC2", bold=True, first=True, spc=1.4)
tf = textbox(s, 7.48, 2.52, 4.90, 2.28)
write(tf, "how much closer this move leaves\nthe gates ready now", 15, WHITE,
      bold=True, font=HEAD, first=True, line_spacing=1.16, space_after=8)
write(tf, "+   ¼  ×   the same, for the gates\n              coming up next",
      15, "9FB3CE", font=HEAD, line_spacing=1.16, space_after=8)
write(tf, "lowest score wins", 12.5, "5FCFC2", italic=True)

card(s, 7.10, 5.08, 5.61, 1.72, fill=PANEL)
tf = textbox(s, 7.48, 5.34, 4.90, 1.24)
write(tf, "The question this paper asks", 14.5, INK, bold=True, font=HEAD,
      first=True, space_after=8)
write(tf, "What does that score have to become when the machine is no longer one "
          "chip, but several?", 14, GREY, line_spacing=1.16)

notes(s, """
SABRE is the incumbent. It was published in 2019 and it is what Qiskit runs by
default, so in practice it is the routing algorithm for superconducting quantum
computers.

The rule is greedy and almost embarrassingly simple, which is exactly why it works at
scale: score each candidate SWAP by how much it reduces the total distance of the
gates that are ready to run, add a discounted lookahead term so you don't get trapped,
take the best one, repeat.

The box on the right is the whole heuristic in words. And the question at the bottom
is this paper's question: that score is designed for one chip. What does it have to
become when the machine is several chips connected by entanglement links?
""")
footer(s, 6)

# ═════════════════════════════════════════════════════════════════════════════
# 7 — Going modular
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Background  ·  5 of 5", "Why quantum computers are going modular")

tf = textbox(s, M, 1.80, 4.10, 4.6)
bullets(tf, [
    ("Yield. ", "The bigger the chip, the likelier some qubit on it is bad — "
     "and the whole chip goes in the bin."),
    ("Crosstalk. ", "Packing more qubits together makes them interfere with each "
     "other more."),
    ("Wiring and cooling. ", "Every qubit needs control lines running into a "
     "dilution refrigerator. There is only so much room."),
    ("So: build small chips and link them.", ""),
], size=14.5, gap=13)

card(s, M, 5.86, 4.10, 1.28, fill=CORAL_TINT)
tf = textbox(s, M + 0.32, 6.06, 4.10 - 0.64, 0.88, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "The links are made of entanglement — and that is the resource this "
          "whole talk is about.", 13.5, "8E322B", bold=True, first=True,
      line_spacing=1.14)

y_img = picture(s, "architectures", x=5.10, w=7.75, y=2.24)
caption(s, 5.10, y_img + 0.26, 7.75,
        "The two devices used in the experiments. Each core is a small chip of 16 "
        "qubits. The teal qubits are communication ports — the only places an "
        "entanglement link can attach, and otherwise ordinary qubits that can hold data.",
        11.5, GREY, h=0.86)

notes(s, """
Nobody is going to build a million-qubit chip. Yield kills you: the larger the die,
the higher the chance that one qubit on it is defective, and then you throw the whole
thing away. Crosstalk gets worse with density. And every single qubit needs physical
control wiring going down into a refrigerator running at ten millikelvin — there is
simply not enough space.

So the roadmaps — IBM's, and essentially everyone else's — go modular: build small,
good chips and connect them.

The picture shows the two machines used in this paper. Call each small chip a "core."
The grey dots are ordinary qubits. The teal dots are communication ports — special
positions that an entanglement link attaches to. Note that a port is still an ordinary
qubit that can hold data when it isn't being used to communicate. That detail turns
out to matter a lot.
""")
footer(s, 7)

# ═════════════════════════════════════════════════════════════════════════════
# 8 — The EPR pair
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Teleportation  ·  1 of 3", "The currency: an EPR pair")

tf = textbox(s, M, 1.80, 5.70, 2.5)
write(tf, "An EPR pair is two qubits, prepared together so that they are entangled, "
          "and then separated — one left in each of two chips.", 15, INK,
      first=True, space_after=12, line_spacing=1.20)
write(tf, "On its own it carries no information. It is a resource: a one-shot ticket "
          "that lets the two chips do something they otherwise could not.",
      15, INK, line_spacing=1.20)

card(s, M, 4.02, 5.70, 2.42, fill=INK)
tf = textbox(s, M + 0.42, 4.30, 5.70 - 0.84, 1.86)
write(tf, "1 EPR pair", 34, WHITE, bold=True, font=HEAD, first=True, space_after=6)
write(tf, "= one qubit moved across one link. A longer trip costs one pair "
          "per hop.", 15, "9FB3CE", line_spacing=1.16, space_after=4)
write(tf, "That is the unit this paper counts, and the unit it is trying to spend "
          "less of.", 13, "5FCFC2", italic=True, line_spacing=1.14)

facts = [("Slow", "Generating a good pair takes far longer than running a local gate. "
                  "The link, not the processor, sets the pace."),
         ("Noisy", "It is the least reliable thing on the machine — pairs often "
                   "have to be made repeatedly and distilled."),
         ("Rate-limited", "A link produces pairs at a fixed rate. You cannot buy your "
                          "way out of a bad routing decision.")]
px, pw = 6.94, 5.77
for i, (t, d) in enumerate(facts):
    y = 1.80 + i * 1.62
    card(s, px, y, pw, 1.42, fill=PANEL)
    tf = textbox(s, px + 0.36, y + 0.24, pw - 0.72, 0.34)
    write(tf, t, 16, TEAL_DARK, bold=True, font=HEAD, first=True)
    tf = textbox(s, px + 0.36, y + 0.64, pw - 0.72, 0.66)
    write(tf, d, 12.5, GREY, first=True, line_spacing=1.15)

caption(s, px, 6.66, pw, "So the compiler's objective: use as few EPR pairs as "
        "possible.", 13, INK, italic=True)

notes(s, """
Now the central object. An EPR pair — named after Einstein, Podolsky and Rosen — is
two qubits prepared together so they're entangled, then separated, with one half
sitting in each of two chips.

By itself it carries no message. Think of it as fuel, or as a ticket: it enables one
operation that the two chips could not otherwise perform, and then it is used up.

The three boxes on the right are why we care. On today's hardware, generating
entanglement between two modules is slow — orders of magnitude slower than a local
gate — it's the noisiest operation on the machine, and the link produces pairs at a
fixed rate you can't exceed.

That's what makes this a well-posed compiler problem: there is one clearly dominant
cost, and the compiler controls how much of it you spend.
""")
footer(s, 8)

# ═════════════════════════════════════════════════════════════════════════════
# 9 — Teleportation step by step
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Teleportation  ·  2 of 3", "How a qubit actually crosses between chips")

tf = textbox(s, M, 1.60, CW, 0.40)
write(tf, "A qubit cannot be copied, and it cannot be put on a wire. Teleportation is "
          "the protocol that moves one anyway — by spending an EPR pair.",
      15, GREY, first=True)

y_img = picture(s, "teleport_protocol", x=0.66, w=12.02, y=2.16)

caps = ["Two entangled qubits are made and separated, one parked in each chip. "
        "This is the fuel, and it is the expensive step.",
        "The sender measures the travelling qubit together with its half of the "
        "pair. The original is destroyed, and two ordinary bits go over a normal wire.",
        "Those two bits say which small correction to apply. The state is now in "
        "core B — moved, never copied."]
for i, c in enumerate(caps):
    caption(s, 0.80 + i * 4.05, y_img + 0.30, 3.62, c, 12, GREY, h=0.92)

card(s, M, 5.94, CW, 0.90, fill=TEAL_TINT)
tf = textbox(s, M + 0.40, 6.12, CW - 0.80, 0.54, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "No faster-than-light anything:  the two classical bits have to arrive "
          "before the state can be rebuilt — and because the original is "
          "destroyed, no copy is ever made.", 14, TEAL_DARK, bold=True, first=True)

notes(s, """
Here is teleportation, which despite the name is an entirely ordinary piece of physics
and is done routinely in laboratories.

The problem it solves: you cannot copy an unknown qubit, and you cannot send one down
a wire. So how do you move a quantum state from chip A to chip B?

Step one: make an EPR pair and separate it, half in each chip. Step two: the sender
performs a joint measurement on the qubit it wants to send, together with its half of
the pair. This destroys the original — that is essential, it's what keeps no-cloning
intact — and produces two ordinary classical bits. Step three: those two bits are sent
over a normal classical wire, and they tell the receiver which of four small
corrections to apply. Apply it, and the state is now sitting in chip B.

Two things people usually ask. No, this is not faster than light: the classical bits
travel at ordinary speed and nothing happens until they arrive. And no, there is never
a moment when two copies exist — the original is gone before the new one appears.

For our purposes, the accounting is what matters: one teleport consumes exactly one
EPR pair, plus whatever local shuffling was needed to get the qubit to a port.
""")
footer(s, 9)

# ═════════════════════════════════════════════════════════════════════════════
# 10 — Teledata vs telegate
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Teleportation  ·  3 of 3", "Two ways to spend a pair")

lw = 6.02
card(s, M, 1.78, lw, 3.14, fill=TEAL_TINT)
tf = textbox(s, M + 0.42, 2.06, lw - 0.84, 0.36)
write(tf, "TELEDATA", 12, TEAL_DARK, bold=True, first=True, spc=1.4)
tf = textbox(s, M + 0.42, 2.50, lw - 0.84, 2.50)
write(tf, "Move the qubit.", 20, INK, bold=True, font=HEAD, first=True,
      space_after=10)
write(tf, "The pair is consumed immediately and the qubit now lives in the other "
          "chip. Every gate it needs after that is local — so one pair can pay "
          "for a long run of work.", 14, INK, line_spacing=1.18, space_after=10)
write(tf, "This is what dSABRE optimises.", 14, TEAL_DARK, bold=True, italic=True)

card(s, M + lw + 0.30, 1.78, lw, 3.14, fill=PANEL)
tf = textbox(s, M + lw + 0.72, 2.06, lw - 0.84, 0.36)
write(tf, "TELEGATE", 12, GREY, bold=True, first=True, spc=1.4)
tf = textbox(s, M + lw + 0.72, 2.50, lw - 0.84, 2.50)
write(tf, "Move the gate.", 20, INK, bold=True, font=HEAD, first=True,
      space_after=10)
write(tf, "Run a two-qubit gate remotely, leaving both qubits at home. One pair can "
          "serve several gates at once — but only if the entanglement survives "
          "that long, and the chip has spare memory to hold it.", 14, INK,
      line_spacing=1.18, space_after=10)
write(tf, "A different problem: it needs gate grouping.", 14, GREY, bold=True,
      italic=True)

card(s, M, 5.20, CW, 1.60, fill=INK)
tf = textbox(s, M + 0.44, 5.44, CW - 0.88, 1.18)
write(tf, "Why teledata, and not both?", 15, "5FCFC2", bold=True, font=HEAD,
      first=True, space_after=7)
write(tf, "Measured on the 64-qubit suite: one teledata move buys that qubit a "
          "residence in its new chip serving 15.75 of its two-qubit gates on average. "
          "An ungrouped telegate serves exactly one. Implemented and allowed to "
          "compete, telegate cost 93.5% more EPR pairs — it only pays in grouped "
          "form, which is a separate research problem.",
      13.5, "C9D2E4", line_spacing=1.18)

notes(s, """
There are two ways to cash in an EPR pair, and the distinction matters for reading the
results later.

Teledata is what I just showed: move the qubit itself. The pair is spent immediately,
and afterwards the qubit is simply in the other chip — everything it does next is
local and free. So one pair can amortise over a long run of subsequent gates.

Telegate does the opposite: it executes a remote two-qubit gate without moving either
qubit. Its attraction is that if you have a whole group of remote gates that fit
together, one entangled state can serve all of them. But that requires holding live
entanglement in memory for the duration of the group, which needs both a spare link
qubit and entanglement that survives that long.

The bottom box is the measurement that justifies the paper's scope. On this benchmark
suite, a single teledata move ends up serving nearly sixteen of that qubit's gates on
average. An ungrouped telegate serves exactly one. We actually implemented telegate
and let it compete — it cost 93.5% more EPR. So telegate only earns its place in
grouped form, which is a genuinely different compilation problem, and we say so
rather than claiming teledata dominates.
""")
footer(s, 10)

# ═════════════════════════════════════════════════════════════════════════════
# 11 — Teleport is not an expensive SWAP
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "The problem", "A teleport is not just an expensive SWAP")

tf = textbox(s, M, 1.56, CW, 0.56)
write(tf, "The tempting shortcut is to reuse SABRE unchanged and simply give "
          "cross-chip edges a bigger weight. That gets the price right and the "
          "physics wrong.", 15, GREY, first=True)

y_img = picture(s, "swap_vs_tele", x=1.07, w=11.20, y=2.26)

card(s, M, 6.26, CW, 0.90, fill=TEAL_TINT)
tf = textbox(s, M + 0.40, 6.44, CW - 0.80, 0.56, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "A SWAP is symmetric and changes nothing about how full a chip is. "
          "A teleport is one-way, needs staging at a port first, and permanently "
          "consumes a seat. Three differences, all invisible to an edge weight.",
      14, TEAL_DARK, bold=True, first=True, line_spacing=1.14)

notes(s, """
This is the conceptual heart of the paper, so let me be slow about it.

The obvious thing to do is take SABRE, which already knows how to score moves by
distance, and just declare that cross-chip edges are expensive — weight ten instead of
weight one. Then a teleport is "a SWAP on a costly edge" and nothing else has to
change.

That is wrong in three specific ways, all on the right-hand panel.

One: a SWAP is symmetric — do it twice and you're back where you started, at the same
price. A teleport is one-way. Getting the qubit back costs a second EPR pair.

Two: a SWAP is legal between any two adjacent qubits. A teleport is legal only from a
communication port — so before you can teleport at all, you may have to shuffle the
qubit across the chip to reach a port, and evict whatever data qubit is currently
parked on it.

Three, and this is the one nobody else models: a SWAP leaves each chip's occupancy
exactly as it was. A teleport permanently moves a qubit from one chip to another, so
the destination has one fewer free seat than before. Do that enough times and the
destination fills up.

An edge weight cannot express any of these. They aren't costs of an edge, they are
consequences for the state of the machine.
""")
footer(s, 11)

# ═════════════════════════════════════════════════════════════════════════════
# 12 — The trap
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "The problem", "And a full chip is a trap")

steps = [("A chip fills up",
          "Every arriving qubit takes a seat. Nothing hands it back, because "
          "teleports are one-way."),
         ("Its port gets occupied",
          "With no free seats, a data qubit ends up parked on the communication "
          "port — the only exit."),
         ("Now nothing can leave",
          "To free the port you must move its occupant somewhere, and there is "
          "nowhere. The route is stuck with gates still waiting.")]
cw_, gap = 3.83, 0.31
for i, (t, d) in enumerate(steps):
    x = M + i * (cw_ + gap)
    card(s, x, 1.78, cw_, 2.70, fill=WHITE, line="DDE3EA")
    circle_num(s, x + 0.34, 2.08, 0.58, i + 1, fill=CORAL)
    tf = textbox(s, x + 0.34, 2.88, cw_ - 0.68, 0.40)
    write(tf, t, 16, INK, bold=True, font=HEAD, first=True)
    tf = textbox(s, x + 0.34, 3.34, cw_ - 0.68, 0.94)
    write(tf, d, 12.5, GREY, first=True, line_spacing=1.16)
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                Inches(x + cw_ + 0.055), Inches(3.02),
                                Inches(0.20), Inches(0.22))
        ar.fill.solid()
        ar.fill.fore_color.rgb = rgb("C9CFD9")
        ar.line.fill.background()
        ar.shadow.inherit = False

card(s, M, 4.86, CW, 1.14, fill=CORAL_TINT)
tf = textbox(s, M + 0.44, 5.08, CW - 0.88, 0.74, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "This is a real failure mode, not a hypothetical. TeleSABRE — the "
          "state-of-the-art router this work compares against — reports the same "
          "problem and leaves a fix to future work. It returns no answer at all on "
          "5 of the 45 instances evaluated here.",
      14, "8E322B", bold=True, first=True, line_spacing=1.16)

tf = textbox(s, M, 6.24, CW, 0.90)
write(tf, "So a distributed router needs two things a single-chip router never "
          "does:  a sense of how full each chip is while it decides, and a way out "
          "when it has painted itself into a corner.", 16, INK, bold=True,
      font=HEAD, first=True, line_spacing=1.18)

notes(s, """
The occupancy issue isn't a footnote. It creates an actual deadlock.

Follow the three boxes. Teleports only ever move qubits one way, so a chip that
receives more than it sends fills up. Once it is full, some data qubit is necessarily
sitting on the communication port, because the port is just an ordinary qubit
position. And now you're stuck: to use the port you must move its occupant, but there
is no free seat to move it into. Gates are still waiting, and no legal move exists.

This is not a hypothetical failure. TeleSABRE, the strongest prior router in this
setting and our main baseline, reports exactly this failure mode in its own paper and
explicitly leaves a fix to future work. In our experiments it returns nothing at all
on five of the forty-five instances.

So the requirement list for a distributed router has two entries a single-chip router
has never needed: it must know how full each chip is while it is choosing, and it must
have a guaranteed way out when it gets stuck.
""")
footer(s, 12)

# ═════════════════════════════════════════════════════════════════════════════
# 13 — the score
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "dSABRE", "Score the whole move, as one thing")

tf = textbox(s, M, 1.58, CW, 0.44)
write(tf, "dSABRE treats “shuffle to a port, evict whoever is on it, teleport” "
          "as a single macro-action and gives it one score. Lower is better.",
      15, GREY, first=True)

blocks = [("what it costs\nto set up",
           "The local SWAPs needed to walk the qubit to a port, plus the SWAPs "
           "needed to clear both ports. An occupied port is priced, not forbidden.",
           TEAL),
          ("what it costs\nin room",
           "A penalty that grows as the destination chip runs low on free seats. "
           "It makes a crowded landing site unattractive without banning it.",
           TEAL),
          ("what it buys",
           "How much closer the move leaves the waiting gate — plus a lookahead "
           "over gates coming up, discounted by how far ahead they are.",
           TEAL)]
cw_, gap = 3.83, 0.31
for i, (t, d, col) in enumerate(blocks):
    x = M + i * (cw_ + gap)
    card(s, x, 2.16, cw_, 2.72, fill=PANEL)
    tf = textbox(s, x + 0.36, 2.42, cw_ - 0.72, 1.00,
                 anchor=MSO_ANCHOR.BOTTOM)
    write(tf, t, 17, INK, bold=True, font=HEAD, first=True, line_spacing=1.12)
    tf = textbox(s, x + 0.36, 3.56, cw_ - 0.72, 1.16)
    write(tf, d, 12.5, GREY, first=True, line_spacing=1.16)
    if i < 2:
        tf = textbox(s, x + cw_ + 0.02, 3.30, 0.27, 0.40)
        write(tf, "+", 20, "B9C0CC", bold=True, first=True, align=PP_ALIGN.CENTER)

card(s, M, 5.10, CW, 1.90, fill=INK)
tf = textbox(s, M + 0.46, 5.32, CW - 0.92, 1.46)
write(tf, "The part that makes it work: capacity enters twice, and the two are "
          "kept apart.", 16.5, WHITE, bold=True, font=HEAD, first=True,
      space_after=10, line_spacing=1.14)
write(tf, "The penalty above ranks candidate moves — it can be outvoted by a "
          "large enough gain. Separately, a hard legality floor decides which moves "
          "are offered at all: no teleport is even generated unless the destination "
          "has at least two free seats. A weight can be outvoted; an admission test "
          "cannot. That separation is what turns a heuristic into something with a "
          "completion guarantee.", 13.5, "C9D2E4", line_spacing=1.20)

notes(s, """
Here is the algorithm's core idea, in three pieces plus one structural decision.

Rather than pricing a cross-chip hop, dSABRE bundles the entire sequence — walk the
qubit to a port, evict whatever data qubit is sitting on the port, teleport — and
scores that whole bundle as one macro-action.

The score has three parts, and they map exactly onto the three differences from the
previous slide. What it costs to set up is the staging and eviction — that's the "you
need a port" problem. What it costs in room is a penalty that grows as the destination
fills — that's the occupancy problem. What it buys is classic SABRE: distance progress
on the waiting gate, plus discounted lookahead.

Now the part I'd emphasise, in the dark box. Capacity appears twice in this design and
the separation is deliberate. The penalty is a soft term that ranks moves — and like
any weight it can be outvoted, if some move offers a large enough gain. So the penalty
alone cannot guarantee anything. Separately there is a hard floor applied when
candidates are generated: a teleport into a chip with fewer than two free seats is
never offered, at any price.

Ranking and legality are different jobs. Keeping them apart is what lets us prove the
router always finishes — and, as the ablation will show, it's also worth a lot in
practice.
""")
footer(s, 13)

# ═════════════════════════════════════════════════════════════════════════════
# 14 — worked example
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "dSABRE", "One decision, worked through")

y_img = picture(s, "teleport_candidates", x=M, w=6.55, y=1.86)
caption(s, M, y_img + 0.18, 6.55, "A gate is pending between q₁ (blue, in C₁) "
        "and q₂ (red, in C₅). Three links leave C₁. Which one?",
        11.5, GREY, h=0.60)

rows = [("", "sets up", "room", "buys", "score"),
        ("A  → C₂, towards q₂", "1", "0", "−12", "−11"),
        ("B  → C₀, away from q₂", "2", "0", "+11", "+13"),
        ("C  → C₄, towards q₂", "3", "15", "−12", "+6")]

tx, ty, tw = 7.52, 1.92, 5.20
tbl = s.shapes.add_table(4, 5, Inches(tx), Inches(ty), Inches(tw),
                         Inches(1.85)).table
tbl.first_row = False
tbl.horz_banding = False
for w_, c_ in zip((2.12, 0.82, 0.72, 0.76, 0.78), range(5)):
    tbl.columns[c_].width = Inches(w_)
for r_ in range(4):
    tbl.rows[r_].height = Inches(0.44 if r_ else 0.36)
    for c_ in range(5):
        cell = tbl.cell(r_, c_)
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        win = (r_ == 1)
        cell.fill.fore_color.rgb = rgb(WHITE if r_ == 0
                                       else (TEAL_TINT if win else PANEL))
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c_ == 0 else PP_ALIGN.CENTER
        run(p, rows[r_][c_],
            11.5 if r_ == 0 else 13,
            GREY if r_ == 0 else (TEAL_DARK if win else INK),
            bold=(r_ == 0 or win))

card(s, tx, 4.02, tw, 2.72, fill=INK)
tf = textbox(s, tx + 0.36, 4.26, tw - 0.72, 2.26)
write(tf, "A wins.", 18, WHITE, bold=True, font=HEAD, first=True, space_after=9)
write(tf, "It has the cheapest staging and the most progress.", 13.5, "C9D2E4",
      space_after=9, line_spacing=1.18)
write(tf, "B moves the qubit the wrong way, so its progress term turns positive and "
          "sinks it.", 13.5, "C9D2E4", space_after=9, line_spacing=1.18)
write(tf, "C buys exactly the same progress as A — but it lands in a nearly-full "
          "chip. The room term alone reverses the decision.", 13.5, "5FCFC2",
      bold=True, line_spacing=1.18)

notes(s, """
Let's watch the score actually decide something, because this example isolates each
term.

Situation: a gate is waiting between q1, sitting in chip C1, and q2, over in C5. Three
entanglement links leave C1, so there are three candidate teleports. The table gives
the three score components for each.

Candidate A goes towards the target and costs one local SWAP to stage. Progress is
minus twelve — remember negative is good — so it scores minus eleven.

Candidate B goes the wrong way. Its progress term is positive eleven, actively worse,
so it scores plus thirteen and is out.

Candidate C is the interesting one. It moves towards the target and buys exactly the
same progress as A — minus twelve, identical. If you scored only by distance, A and C
would be nearly tied. But C lands in a chip with only two free seats, so the room term
adds fifteen, and it also has an occupied port to evacuate. Final score: plus six.

The capacity term, on its own, flips the decision. That's the mechanism working — and
the ablation later shows it's worth about nine percent across the whole suite.
""")
footer(s, 14)

# ═════════════════════════════════════════════════════════════════════════════
# 15 — the loop
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "dSABRE", "The routing loop")

items = [("Run everything that's ready.",
          "Drain the front layer of every gate that can already execute."),
         ("Sort what's left.",
          "Same-chip gates that just need shuffling, versus gates whose qubits are "
          "on different chips."),
         ("Local work first.",
          "A SWAP is always cheaper than an EPR pair, so no teleport is even scored "
          "while same-chip work remains. Only then: score the teleports and take the "
          "best legal one."),
         ("Checkpoint, or recover.",
          "Every time the circuit gets shorter, save the state. If nothing shrinks "
          "for too long, rewind to the last checkpoint and call recovery.")]
tf = textbox(s, M, 1.82, 4.55, 4.9)
for i, (t, d) in enumerate(items):
    p = para(tf, i == 0)
    p.space_after = Pt(4)
    run(p, "P%d   " % (i + 1), 13, TEAL, bold=True, font=HEAD)
    run(p, t, 14.5, INK, bold=True)
    p = para(tf)
    p.space_after = Pt(13)
    p.line_spacing = 1.16
    run(p, d, 13, GREY)

y_img = picture(s, "flowchart_talk", x=5.42, w=7.42, y=2.02)
caption(s, 5.42, y_img + 0.24, 7.42, "The loop as implemented. Everything on the right "
        "of the diagram is bookkeeping to make sure the router can always back out of "
        "a bad state.", 11.5, GREY, h=0.62)

notes(s, """
The loop itself, in four steps.

P1: run everything that can already run. Free progress, always take it first.

P2: look at what's left and split it — gates whose two qubits are on the same chip and
just need shuffling, versus gates whose qubits are on different chips.

P3 is a priority rule worth pausing on. dSABRE will not even score a teleport while
any same-chip gate is still waiting. The reasoning: a SWAP is unconditionally cheaper
than an EPR pair, so spending a pair while local work is outstanding means you've
imposed some arbitrary exchange rate between the two. We measured the baseline here:
under TeleSABRE's pooled scoring, 59.5% of its cross-chip operations fire while a
same-chip gate is still waiting.

P4 is the safety net. Every time the remaining circuit gets smaller, we snapshot the
state. If enough iterations go by with no progress at all, we roll back to that
snapshot and invoke the recovery procedure — which is the next slide.
""")
footer(s, 15)

# ═════════════════════════════════════════════════════════════════════════════
# 16 — the guarantee
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "dSABRE", "Three rules that make getting stuck impossible")

rules = [("Always leave two seats",
          "A teleport is offered only if the destination has at least two free "
          "seats. One receives the arrival; the other keeps the chip from ever "
          "becoming full — so a port can always be cleared."),
         ("Start with room",
          "The initial placement reserves spare seats in every chip before routing "
          "begins, so the condition holds from the first step rather than being "
          "hoped for."),
         ("Keep a fire exit",
          "If the ordinary loop stalls anyway, a bounded recovery transaction walks "
          "a free seat over from a chip that has one to spare, brings the two "
          "waiting qubits together, and runs the gate.")]
tf = textbox(s, M, 1.80, 6.95, 4.6)
for i, (t, d) in enumerate(rules):
    y = 1.80 + i * 1.62
    circle_num(s, M, y + 0.02, 0.52, i + 1)
    tf2 = textbox(s, M + 0.76, y + 0.02, 6.19, 0.40)
    write(tf2, t, 17, INK, bold=True, font=HEAD, first=True)
    tf2 = textbox(s, M + 0.76, y + 0.48, 6.19, 1.00)
    write(tf2, d, 13, GREY, first=True, line_spacing=1.18)

card(s, 8.02, 1.80, 4.70, 2.56, fill=INK)
tf = textbox(s, 8.38, 2.06, 3.98, 2.10)
write(tf, "THEOREM", 10.5, "5FCFC2", bold=True, first=True, spc=1.4,
      space_after=9)
write(tf, "Under conditions you can check before compiling, dSABRE routes every gate "
          "and never aborts.", 16, WHITE, bold=True, font=HEAD,
      line_spacing=1.16, space_after=9)
write(tf, "Connected chips, at least four sites each, and a little global slack. "
          "The tightest case in the paper has 32 spare seats where 13 suffice.",
      12.5, "9FB3CE", line_spacing=1.16)

card(s, 8.02, 4.56, 4.70, 2.10, fill=TEAL_TINT)
tf = textbox(s, 8.38, 4.82, 3.98, 1.62)
write(tf, "In practice it's a fire exit", 15, TEAL_DARK, bold=True, font=HEAD,
      first=True, space_after=8)
write(tf, "No reported route at 25–64 qubits uses recovery at all. At 100, 200 "
          "and 360 qubits it retires 6, 10 and 66 gates. Across 1074 transactions, "
          "no precondition ever failed.", 12.5, GREY, line_spacing=1.18)

notes(s, """
This is the completion guarantee, and it comes from three rules rather than anything
clever.

Rule one is the legality floor. A teleport is generated only if the destination has at
least two free seats. Why two? One is consumed by the arrival. The second guarantees
the chip is never left completely full, which in turn guarantees that its
communication port can always be cleared later. That's the deadlock from earlier,
closed off at the source.

Rule two makes the condition true at the start: the initial placement deliberately
holds back seats in every chip, so we begin in a good state rather than hoping for one.

Rule three is the fallback. If the loop stalls anyway, a bounded transaction relays a
free seat over from a chip that has one to spare, walks one of the two waiting qubits
to meet the other, and executes the gate. It's transactional: it either completes or
leaves the state untouched.

Put together, the theorem says dSABRE routes every gate and never aborts, under
conditions you can check before you start compiling. The conditions are loose — the
tightest instance in the paper has thirty-two spare seats where thirteen would do.

And I want to be honest about what this buys in practice, in the teal box: on the
small and medium instances recovery never fires at all. It's insurance. It starts
mattering at scale — sixty-six gates at three hundred and sixty qubits.
""")
footer(s, 16)

# ═════════════════════════════════════════════════════════════════════════════
# 17 — headline result
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Results  ·  1 of 4", "About half the entanglement, everywhere we looked")

tf = textbox(s, M, 1.58, CW, 0.40)
write(tf, "Against TeleSABRE, on the same devices, the same 21 benchmark circuits, "
          "and the same run budget — best of three seeds for both routers.",
      14.5, GREY, first=True)

y_img = picture(s, "chart_headline", x=1.47, w=10.40, y=2.06)

card(s, M, 6.34, CW, 0.86, fill=TEAL_TINT)
tf = textbox(s, M + 0.40, 6.52, CW - 0.80, 0.52, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "Two of these devices are networks of IBM heavy-hex chips — irregular, "
          "with dangling qubits and a hub every route must cross. Nothing in dSABRE "
          "is tuned for them, and the result holds.", 14, TEAL_DARK, bold=True,
      first=True, line_spacing=1.14)

notes(s, """
The headline. Across five different device-and-workload settings, dSABRE uses between
forty-one and sixty-two percent fewer EPR pairs than TeleSABRE, the strongest prior
router in this setting.

A few words on why this is a fair comparison, because "we beat the baseline" claims in
this area are often not. TeleSABRE uses the same communication primitive we do —
teledata — counts local SWAPs the same way, and we run it on its own benchmark
topologies. Both routers report best-of-three seeds, which is standard use for both.
And if you take the median seed instead of the best, the reductions are fifty, sixty
and forty-three percent — so this isn't an artefact of the best-of-N convention.

The bottom two bars are worth pointing out. Those are networks of IBM heavy-hex chips
— completely irregular, with dangling leaf qubits and, in the star case, a hub that
every cross-chip route has to pass through. Nothing in the algorithm is specialised
for them; the seat-reservation logic just reads vertex degree, so it picks leaves
where it picked grid corners. The result carries over.
""")
footer(s, 17)

# ═════════════════════════════════════════════════════════════════════════════
# 18 — per-circuit
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Results  ·  2 of 4", "Not an artefact of averaging")

tf = textbox(s, M, 1.58, CW, 0.40)
write(tf, "Every circuit in the 64-qubit suite, both routers, log scale. dSABRE wins "
          "on all of them.", 14.5, GREY, first=True)

y_img = picture(s, "chart_percircuit", x=1.42, w=10.50, y=2.02)

card(s, M, 6.30, CW, 0.88, fill=CORAL_TINT)
tf = textbox(s, M + 0.40, 6.48, CW - 0.80, 0.54, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "The last column is the sharper result: on Random, TeleSABRE returns no "
          "routing at all in any of three attempts. dSABRE completes it with 732 "
          "pairs. Fewer pairs is an average; finishing is not.",
      14, "8E322B", bold=True, first=True, line_spacing=1.14)

notes(s, """
Averages can hide a lot, so here is every circuit in the largest suite, both routers,
on a log scale because the circuits span three orders of magnitude in size.

dSABRE is lower on every single circuit both routers complete. The margin varies —
biggest on the structured circuits like GHZ and Graphstate where good placement pays
off enormously, smallest on Multiplier, which at thirteen thousand CX gates is dense
enough that there isn't much room to be clever.

The rightmost column is the one I'd actually highlight. On the Random circuit,
TeleSABRE returns nothing — all three seeds exhaust its internal attempt limit at its
own deadlock criterion, well inside the time budget. That's its own verdict, not a
timeout we imposed. dSABRE routes it with 732 pairs.

That distinction matters for a compiler. Using fewer pairs is a quantitative
improvement you can average. Returning an answer at all is not something you can
average — you either compile the program or you don't.
""")
footer(s, 18)

# ═════════════════════════════════════════════════════════════════════════════
# 19 — scale and completion
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Results  ·  3 of 4", "It keeps working as the machine grows")

y_img = picture(s, "chart_scale", x=1.27, w=10.80, y=1.54)

stats = [("48 s", "to compile a 360-qubit QFT across twenty chips."),
         ("1.9 s vs 4.7 s", "geometric-mean compile time at 64 qubits — dSABRE "
                            "in Python, the baseline a compiled C++ binary."),
         ("≈ linear", "growth in compile time with gate count at a fixed "
                       "device: fitted exponent 0.97, R² = 0.95.")]
cw_, gap = 3.83, 0.31
for i, (big, d) in enumerate(stats):
    x = M + i * (cw_ + gap)
    card(s, x, 5.60, cw_, 1.36, fill=PANEL)
    tf = textbox(s, x + 0.32, 5.76, cw_ - 0.64, 0.42)
    write(tf, big, 21, TEAL_DARK, bold=True, font=HEAD, first=True)
    tf = textbox(s, x + 0.32, 6.20, cw_ - 0.64, 0.66)
    write(tf, d, 11.5, GREY, first=True, line_spacing=1.14)

notes(s, """
Scaling. We take two circuit families and grow them — a hundred, two hundred, three
hundred and sixty logical qubits, on six, twelve and twenty chips respectively, so
that occupancy stays roughly constant.

The dashed red boxes are where TeleSABRE returns nothing. It manages both families at
a hundred qubits, one of them at two hundred, and neither at three hundred and sixty.
dSABRE completes all six.

I want to be careful here: this is evidence that dSABRE keeps working, not a claim
about why the baseline stops. Its non-convergence could be search behaviour or an
implementation limit, and we don't attribute it to capacity handling.

The right panel is completion over the whole evaluation: forty-five out of forty-five
against forty.

And the compile times at the bottom. Forty-eight seconds for a three-hundred-and-sixty
qubit QFT. At sixty-four qubits, a geometric mean of 1.9 seconds against the
baseline's 4.7 — and I should stress the handicap there runs the wrong way: dSABRE is
pure Python and TeleSABRE is compiled C++. Compile time grows very close to linearly
in gate count for a fixed device — the fitted exponent is 0.97.
""")
footer(s, 19)

# ═════════════════════════════════════════════════════════════════════════════
# 20 — ablation
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Results  ·  4 of 4", "Which parts are actually doing the work?")

tf = textbox(s, M, 1.58, CW, 0.40)
write(tf, "Remove one component at a time and re-run the whole 64-qubit suite. "
          "Longer bar = the component was doing more.", 14.5, GREY, first=True)

y_img = picture(s, "chart_ablation", x=1.47, w=10.40, y=2.02)

card(s, M, 6.28, CW, 0.90, fill=INK)
tf = textbox(s, M + 0.42, 6.46, CW - 0.84, 0.56, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "The two capacity mechanisms are substitutes, not separate contributions: "
          "either one alone is nearly free, but remove both and the cost almost "
          "doubles and two circuits stop finishing. +0.2% and +9.4% become +87.8%.",
      13.5, "C9D2E4", bold=True, first=True, line_spacing=1.16)

notes(s, """
Ablations, because "our system is better" is not an explanation.

The depth-ordered lookahead — expanding the set of upcoming gates breadth-first rather
than in topological order, so that "how far ahead" actually means something — is worth
about twelve percent.

The capacity penalty is worth about nine.

Now look at the interesting pair. Removing the hard legality floor on its own costs
almost nothing: two-tenths of a percent. Removing the capacity penalty on its own
costs nine. But removing both costs eighty-eight percent, and two circuits stop
finishing entirely.

That's not additive, and the reason is that the two mechanisms are substitutes. The
penalised band — destination running low on seats — contains the forbidden band. So
while the soft penalty is present, the router rarely proposes an illegal landing in
the first place, and the floor has little left to do.

I want to flag what that does and does not license. It does not mean you can drop the
floor. A penalty, however heavy, can be outvoted by a large enough progress term; an
admission test cannot. The theorem needs the floor on every teleport. What the number
shows is that on these instances the mechanisms cover for each other — and that
capacity handling as a whole is the single most valuable thing in the design.
""")
footer(s, 20)

# ═════════════════════════════════════════════════════════════════════════════
# 21 — honest reading
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide()
head(s, "Caveats", "How to read these numbers honestly")

lw = 6.02
card(s, M, 1.78, lw, 3.98, fill=PANEL)
tf = textbox(s, M + 0.40, 2.04, lw - 0.80, 0.34)
write(tf, "AGAINST pytket-dqc", 11, GREY, bold=True, first=True, spc=1.4)
tf = textbox(s, M + 0.40, 2.48, lw - 0.80, 3.40)
write(tf, "A different cost model, not a different router.", 16, INK, bold=True,
      font=HEAD, first=True, space_after=10, line_spacing=1.14)
write(tf, "It uses gate teleportation, so one pair can serve many gates. On its own "
          "terms dSABRE wins at 25 and 36 qubits (−56%, −14%) and loses at 64, "
          "where dSABRE uses 15% more.", 13.5, INK, line_spacing=1.18,
      space_after=10)
write(tf, "But its distributions assume unlimited link qubits: 4 of 6, 2 of 6, and at "
          "least 6 of 9 of them provably cannot be built on the device we evaluate. "
          "And its count assumes entanglement surviving more than ten gates. Cap it "
          "at three and the 64-qubit result flips to 59.6% in dSABRE's favour.",
      13.5, GREY, line_spacing=1.18)

card(s, M + lw + 0.30, 1.78, lw, 3.98, fill=CORAL_TINT)
tf = textbox(s, M + lw + 0.70, 2.04, lw - 0.80, 0.34)
write(tf, "WHAT THIS DOES NOT CLAIM", 11, "8E322B", bold=True, first=True, spc=1.4)
tf = textbox(s, M + lw + 0.70, 2.48, lw - 0.80, 3.40)
write(tf, "One resource, one communication model.", 16, INK, bold=True, font=HEAD,
      first=True, space_after=10, line_spacing=1.14)
bullets(tf, [
    "We minimise EPR pairs, not runtime and not output fidelity.",
    "The trace we emit is unscheduled — so fewer pairs does not automatically "
    "mean a shorter or more accurate run.",
    "Teledata is not claimed to dominate telegate; grouped telegate is a separate "
    "problem we did not solve.",
    "Latency, coherence time and link contention are outside the model.",
], size=13.5, gap=8, bullet_color="C0453B", first=False)

tf = textbox(s, M, 6.06, CW, 0.90)
write(tf, "The controlled TeleSABRE comparison is the result. The cross-model studies "
          "are context — they say where a different primitive would win, and "
          "under what assumptions.", 15, INK, bold=True, font=HEAD, first=True,
      line_spacing=1.16)

notes(s, """
I want to spend a moment here rather than skip to the conclusion, because comparing
distributed quantum compilers is genuinely treacherous — two tools can both report
"EPR pairs" and be counting different things.

pytket-dqc is the main example. It's an excellent tool, but it uses gate teleportation
with hypergraph partitioning, so one pair can serve a whole group of gates. On its own
terms we beat it comfortably at twenty-five and thirty-six qubits and it beats us by
fifteen percent at sixty-four.

Two assumptions sit underneath that number. First, its input network allows unbounded
live link qubits per chip — and when we checked, four of six, two of six, and at least
six of nine of its distributions provably cannot be materialised within the
communication capacity the device actually has. One twenty-five-qubit case wants twenty
simultaneous link qubits at a chip that provides two. Second, its count assumes
entanglement surviving more than ten gates. Re-score the same distributions with a cap
of three gates per pair and the sixty-four qubit result flips to a sixty percent
advantage for us.

That's not a gotcha — it's the honest statement that the comparison depends on what
hardware you assume.

And the right-hand column is what we don't claim. We minimise one resource. The output
isn't scheduled, so a lower EPR count doesn't automatically give a shorter or more
accurate run. That's the next piece of work, not a result we're hiding.
""")
footer(s, 21)

# ═════════════════════════════════════════════════════════════════════════════
# 22 — takeaways
# ═════════════════════════════════════════════════════════════════════════════
s = new_slide(dark=True)
head(s, "Wrapping up", "Three things to take away", dark=True)

takes = [("Price the whole move, not the hop.",
          "A cross-chip transfer is a macro-action: staging to a port, evicting "
          "whoever sits there, and the seat it permanently consumes. An edge weight "
          "cannot express any of that."),
         ("Keep ranking separate from legality.",
          "A penalty can be outvoted; an admission test cannot. That separation is "
          "what turns a greedy heuristic into a router with a completion guarantee."),
         ("It pays, and it scales.",
          "49–62% fewer EPR pairs than the state of the art on the main suites, "
          "45 of 45 instances routed against 40, and a 360-qubit circuit compiled "
          "in 48 seconds.")]
for i, (t, d) in enumerate(takes):
    y = 1.80 + i * 1.38
    circle_num(s, M, y, 0.54, i + 1, fill=TEAL, color=WHITE)
    tf = textbox(s, M + 0.80, y + 0.01, 7.80, 0.38)
    write(tf, t, 17.5, WHITE, bold=True, font=HEAD, first=True)
    tf = textbox(s, M + 0.80, y + 0.46, 7.80, 0.82)
    write(tf, d, 13, "9FB3CE", first=True, line_spacing=1.18)

card(s, M, 5.90, 8.60, 1.06, fill="17233E")
tf = textbox(s, M + 0.40, 6.10, 8.60 - 0.80, 0.68, anchor=MSO_ANCHOR.MIDDLE)
write(tf, "Next:  choosing jointly between moving the qubit and moving the gate, "
          "then scheduling the result for time and fidelity.", 13.5, "C9D2E4",
      bold=True, first=True, line_spacing=1.14)

card(s, 9.62, 1.86, 3.10, 5.18, fill="17233E")
tf = textbox(s, 9.94, 2.20, 2.46, 4.50)
write(tf, "CODE & DATA", 10.5, "5FCFC2", bold=True, first=True, spc=1.4,
      space_after=10)
write(tf, "github.com\n/ebony72\n/dsabre", 15, WHITE, bold=True, font=HEAD,
      line_spacing=1.22, space_after=14)
write(tf, "Reference implementation, all 45 benchmark circuits, per-circuit result "
          "files, and the online appendices with the full proof.", 12,
      "9FB3CE", line_spacing=1.18, space_after=14)
write(tf, "Thank you — questions?", 13.5, "5FCFC2", bold=True, italic=True,
      line_spacing=1.16)

notes(s, """
Three things to take away.

First, the design principle: price the whole move rather than the hop. A cross-chip
transfer is not a SWAP on an expensive edge — it is a macro-action with a setup cost,
a direction, and a permanent consequence for how full the destination is. Model it as
one thing and score it as one thing.

Second, the structural point, which I think generalises past this paper: keep ranking
separate from legality. Soft penalties rank; hard admission tests decide what is
allowed. Conflating them is what makes greedy routers fail unpredictably at
saturation, and separating them is what let us prove this one always finishes.

Third, it pays: roughly half the entanglement of the state of the art, every instance
routed, and it scales to hundreds of qubits in tens of seconds.

What's next is joint selection — deciding per-move whether to move the qubit or move
the gate — and then scheduling the result, because EPR count is not the same as
wall-clock time or fidelity.

Everything is on GitHub: the implementation, the benchmark circuits, the per-circuit
results, and the appendices with the full proof.

Thank you — happy to take questions.
""")
footer(s, 22)

out = os.path.join(_HERE, "dSABRE_talk.pptx")
prs.save(out)
print("wrote", out, "with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
