#!/usr/bin/env python3
"""Render dSABRE_talk.pptx to PNGs for visual QA, without LibreOffice.

Reads the saved .pptx back through python-pptx and repaints every shape with
PIL, using the real Calibri/Cambria font files PowerPoint itself ships, so
text wrapping and overflow are measured with true metrics.

Also prints a report of shapes that overflow their box or leave the slide.

Run:  python3 preview.py            # writes preview/slide-NN.png
"""

import os, sys, glob
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

_HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(_HERE, "preview")
DPI = 110
DF = "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts"

FONTS = {
    ("Calibri", False, False): f"{DF}/Calibri.ttf",
    ("Calibri", True,  False): f"{DF}/Calibrib.ttf",
    ("Calibri", False, True):  f"{DF}/Calibrii.ttf",
    ("Calibri", True,  True):  f"{DF}/Calibriz.ttf",
    ("Cambria", False, False): f"{DF}/Cambria.ttc",
    ("Cambria", True,  False): f"{DF}/Cambriab.ttf",
    ("Cambria", False, True):  f"{DF}/Cambriai.ttf",
    ("Cambria", True,  True):  f"{DF}/Cambriaz.ttf",
    ("Arial",   False, False): "/System/Library/Fonts/Supplemental/Arial.ttf",
    ("Arial",   True,  False): "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
}
_cache = {}


def font(name, size_pt, bold=False, italic=False):
    key = (name, bold, italic, round(size_pt, 1))
    if key not in _cache:
        path = FONTS.get((name, bold, italic)) or FONTS[("Calibri", bold, italic)]
        px = max(1, int(round(size_pt * DPI / 72.0)))
        _cache[key] = ImageFont.truetype(path, px)
    return _cache[key]


def px(emu):
    return emu * DPI / 914400.0


def hexcol(c, default=(20, 24, 40)):
    try:
        return tuple(bytes.fromhex(str(c)))
    except Exception:
        return default


issues = []


# ── text layout ──────────────────────────────────────────────────────────────
def para_runs(p):
    out = []
    for r in p.runs:
        f = r.font
        out.append({
            "t": r.text,
            "sz": f.size.pt if f.size else 18.0,
            "b": bool(f.bold),
            "i": bool(f.italic),
            "n": f.name or "Calibri",
            "c": hexcol(f.color.rgb) if (f.color and f.color.type is not None
                                         and getattr(f.color, "rgb", None)) else (20, 24, 40),
        })
    return out


def bullet_of(p):
    pPr = p._p.find(qn("a:pPr"))
    if pPr is None:
        return None
    ch = pPr.find(qn("a:buChar"))
    if ch is None:
        return None
    clr = pPr.find(qn("a:buClr"))
    col = (18, 134, 126)
    if clr is not None:
        s = clr.find(qn("a:srgbClr"))
        if s is not None:
            col = hexcol(s.get("val"))
    marL = int(pPr.get("marL") or 0)
    return {"char": ch.get("char", "•"), "color": col, "marL": px(marL)}


def wrap(draw, runs, width):
    """Greedy wrap across runs -> list of lines; each line = [(text, run)]."""
    lines, cur, cw = [], [], 0.0
    for r in runs:
        fnt = font(r["n"], r["sz"], r["b"], r["i"])
        for chunk in r["t"].split("\n"):
            if chunk is not r["t"].split("\n")[0]:
                lines.append(cur); cur, cw = [], 0.0
            words = chunk.split(" ")
            for wi, w in enumerate(words):
                seg = (" " if (cur and wi >= 0 and cw > 0) else "") + w
                sw = draw.textlength(seg, font=fnt)
                if cw + sw > width and cur:
                    lines.append(cur); cur, cw = [], 0.0
                    seg, sw = w, draw.textlength(w, font=fnt)
                cur.append((seg, r)); cw += sw
    lines.append(cur)
    return lines


def draw_tf(draw, tf, x, y, w, h, label=""):
    """Render a text frame; returns used height. Records overflow."""
    lines_all = []
    for p in tf.paragraphs:
        runs = para_runs(p)
        if not runs:
            lines_all.append(("gap", p, []))
            continue
        bu = bullet_of(p)
        ind = bu["marL"] if bu else 0.0
        ls = p.line_spacing or 1.0
        sa = p.space_after.pt if p.space_after else 0.0
        lines = wrap(draw, runs, w - ind)
        lines_all.append(("p", (p, bu, ind, ls, sa, runs), lines))

    total = 0.0
    for kind, meta, lines in lines_all:
        if kind == "gap":
            continue
        p, bu, ind, ls, sa, runs = meta
        lh = max(r["sz"] for r in runs) * 1.20 * DPI / 72.0 * ls
        total += lh * len(lines) + sa * DPI / 72.0

    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        cy = y + (h - total) / 2.0
    elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
        cy = y + max(0.0, h - total)
    else:
        cy = y

    for kind, meta, lines in lines_all:
        if kind == "gap":
            continue
        p, bu, ind, ls, sa, runs = meta
        lh = max(r["sz"] for r in runs) * 1.20 * DPI / 72.0 * ls
        for li, line in enumerate(lines):
            lw = sum(draw.textlength(t, font=font(r["n"], r["sz"], r["b"], r["i"]))
                     for t, r in line)
            if p.alignment == PP_ALIGN.CENTER:
                cx = x + ind + (w - ind - lw) / 2.0
            elif p.alignment == PP_ALIGN.RIGHT:
                cx = x + w - lw
            else:
                cx = x + ind
            if bu and li == 0:
                bf = font("Arial", runs[0]["sz"] * 0.95, False, False)
                draw.text((x, cy + lh * 0.06), bu["char"], font=bf, fill=bu["color"])
            for t, r in line:
                fnt = font(r["n"], r["sz"], r["b"], r["i"])
                draw.text((cx, cy + lh * 0.06), t, font=fnt, fill=r["c"])
                cx += draw.textlength(t, font=fnt)
            cy += lh
        cy += sa * DPI / 72.0

    if total > h + 2:
        issues.append((label, "TEXT OVERFLOW: needs %.2f\" in a %.2f\" box  |  %s"
                       % (total / DPI, h / DPI,
                          (tf.text or "")[:64].replace("\n", " "))))
    return total


# ── shape painting ───────────────────────────────────────────────────────────
def paint(draw, sl, shp, sn, imgobj):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    x, y = px(shp.left or 0), px(shp.top or 0)
    w, h = px(shp.width or 0), px(shp.height or 0)
    label = "slide %d" % sn

    if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            im = Image.open(shp.image.blob and __import__("io").BytesIO(shp.image.blob))
            im = im.convert("RGBA").resize((max(1, int(w)), max(1, int(h))))
            imgobj.paste(im, (int(x), int(y)), im)
        except Exception as e:
            draw.rectangle([x, y, x + w, y + h], outline=(200, 60, 60), width=2)
        return

    if shp.shape_type == MSO_SHAPE_TYPE.TABLE:
        tbl = shp.table
        cy = y
        for r_ in range(len(tbl.rows)):
            rh = px(tbl.rows[r_].height)
            cx = x
            for c_ in range(len(tbl.columns)):
                cwid = px(tbl.columns[c_].width)
                cell = tbl.cell(r_, c_)
                try:
                    fc = hexcol(cell.fill.fore_color.rgb, (245, 247, 249))
                except Exception:
                    fc = (245, 247, 249)
                draw.rectangle([cx, cy, cx + cwid, cy + rh], fill=fc)
                draw_tf(draw, cell.text_frame, cx + px(cell.margin_left),
                        cy + px(cell.margin_top),
                        cwid - px(cell.margin_left) - px(cell.margin_right),
                        rh - px(cell.margin_top) - px(cell.margin_bottom),
                        label + " table")
                cx += cwid
            cy += rh
        return

    if shp.shape_type == MSO_SHAPE_TYPE.LINE or shp.element.tag.endswith("cxnSp"):
        try:
            col = hexcol(shp.line.color.rgb, (120, 130, 150))
        except Exception:
            col = (120, 130, 150)
        draw.line([x, y, x + w, y + h], fill=col, width=3)
        return

    is_auto = shp.element.tag.endswith("}sp")
    if is_auto and shp.has_text_frame:
        try:
            has_fill = shp.fill.type is not None and shp.fill.type == 1
        except Exception:
            has_fill = False
        fill = None
        if has_fill:
            try:
                fill = hexcol(shp.fill.fore_color.rgb, None)
            except Exception:
                fill = None
        outline = None
        try:
            if shp.line.fill.type == 1:
                outline = hexcol(shp.line.color.rgb, None)
        except Exception:
            pass

        geom = shp.element.find(".//" + qn("a:prstGeom"))
        nm = (geom.get("prst") if geom is not None else "rect") or "rect"
        if fill or outline:
            if nm in ("ellipse",):
                draw.ellipse([x, y, x + w, y + h], fill=fill, outline=outline, width=2)
            elif nm == "roundRect":
                draw.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) * 0.11,
                                       fill=fill, outline=outline, width=2)
            elif nm == "rightArrow":
                draw.polygon([(x, y + h * 0.30), (x + w * 0.55, y + h * 0.30),
                              (x + w * 0.55, y), (x + w, y + h / 2),
                              (x + w * 0.55, y + h), (x + w * 0.55, y + h * 0.70),
                              (x, y + h * 0.70)], fill=fill)
            else:
                draw.rectangle([x, y, x + w, y + h], fill=fill, outline=outline, width=2)
        tf = shp.text_frame
        if (tf.text or "").strip():
            draw_tf(draw, tf, x + px(tf.margin_left), y + px(tf.margin_top),
                    w - px(tf.margin_left) - px(tf.margin_right),
                    h - px(tf.margin_top) - px(tf.margin_bottom), label)
        return


def main():
    prs = Presentation(os.path.join(_HERE, "dSABRE_talk.pptx"))
    W = int(px(prs.slide_width)); H = int(px(prs.slide_height))
    os.makedirs(OUT, exist_ok=True)
    for f in glob.glob(os.path.join(OUT, "*.png")):
        os.remove(f)

    for i, sl in enumerate(prs.slides, 1):
        try:
            bg = hexcol(sl.background.fill.fore_color.rgb, (255, 255, 255))
        except Exception:
            bg = (255, 255, 255)
        im = Image.new("RGB", (W, H), bg)
        d = ImageDraw.Draw(im)
        for shp in sl.shapes:
            l, t = px(shp.left or 0), px(shp.top or 0)
            r_, b_ = l + px(shp.width or 0), t + px(shp.height or 0)
            if l < -2 or t < -2 or r_ > W + 2 or b_ > H + 2:
                issues.append(("slide %d" % i,
                               "OFF-SLIDE: %s at (%.2f,%.2f)-(%.2f,%.2f)\""
                               % (shp.shape_type, l / DPI, t / DPI, r_ / DPI, b_ / DPI)))
            paint(d, sl, shp, i, im)
            d = ImageDraw.Draw(im)
        # pictures must not intersect any filled card
        boxes = []
        for shp in sl.shapes:
            geom = shp.element.find(".//" + qn("a:prstGeom"))
            kind = ("pic" if shp.shape_type == 13 else
                    ("card" if (geom is not None and geom.get("prst") == "roundRect"
                                and (shp.width or 0) > 914400) else None))
            if kind:
                boxes.append((kind, shp, px(shp.left or 0), px(shp.top or 0),
                              px(shp.width or 0), px(shp.height or 0)))
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                ka, _, ax, ay, aw, ah = boxes[a]
                kb, _, bx, by, bw, bh = boxes[b]
                if ka == kb == "card":
                    continue
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 2 and oy > 2:
                    issues.append(("slide %d" % i,
                        "OVERLAP: %s and %s overlap by %.2f x %.2f\""
                        % (ka, kb, ox / DPI, oy / DPI)))
        im.save(os.path.join(OUT, "slide-%02d.png" % i))

    print("rendered %d slides -> %s" % (len(prs.slides._sldIdLst), OUT))
    if issues:
        print("\n%d issue(s):" % len(issues))
        for where, what in issues:
            print("  %-10s %s" % (where, what))
    else:
        print("\nno geometry or overflow issues detected")


if __name__ == "__main__":
    main()
